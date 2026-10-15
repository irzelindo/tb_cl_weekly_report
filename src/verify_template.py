"""Validar se o template local corresponde à versão TRL multilaboratório."""
import os
from pathlib import Path
from openpyxl import load_workbook

BASE = Path(__file__).resolve().parent.parent
expected = ["Resumo", "Comparação", "TRL", "Pendências", "Resistência",
            "Resultados", "Desagregação Sexo"]
p = Path(os.getenv("REPORT_TEMPLATE", "templates/TB_Weekly_Report_Template.xlsx"))
if not p.is_absolute():
    p = BASE / p
if not p.is_file():
    raise SystemExit(f"ERRO: ficheiro ausente: {p}")
wb = load_workbook(p, read_only=False, data_only=False)
if wb.sheetnames != expected:
    raise SystemExit(f"ERRO: folhas não coincidem:\n{wb.sheetnames}\nEsperado:\n{expected}")
assert wb["Pendências"]["E4"].value == "45–60 dias", (
    "Intervalo 44–60 sobrepõe 8–44; usar 45–60 dias."
)
assert wb["TRL"]["B21"].value == "TSA 1 e 2", (
    "Falta o quinto indicador TRL (TSA 1 e 2) na linha 21."
)
assert wb["Resultados"]["B2"].value == "Código do DISA"
assert wb["Resultados"]["C2"].value == "Nome do Paciente"
assert wb["Desagregação Sexo"]["D4"].value == "Homens"
assert wb["Desagregação Sexo"]["E4"].value == "Mulheres"
print("PASSOU: template TRL, 7 folhas, faixas sem sobreposição, sexo e TSA.")
wb.close()
