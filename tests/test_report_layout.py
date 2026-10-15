"""Regressão da apresentação final — sem SQL Server nem dados pessoais.

Executar na raiz do projecto:
    python -m unittest discover -s tests -p 'test_report_layout.py'
"""
import unittest
from datetime import date
from openpyxl import Workbook

from src.multilab_writer import (
    LABS, remove_empty_summary_alerts, write_rejections,
)


class ReportPresentationTests(unittest.TestCase):
    def test_summary_hides_empty_alerts_and_preserves_legend(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Resumo"
        ws.merge_cells("A41:F41")
        ws["A41"] = "3. LEITURA RÁPIDA / ALERTAS"
        ws["A42"] = "Alerta"
        ws["A43"] = "Pendências elevadas"
        original = [
            "DESAGREGAÇÃO POR SEXO",
            "LEGENDA",
            "CMTB = Complexo Mycobacterium tuberculosis",
            "RR/MR, Pré-XR e XR",
            "Pendentes = amostras registadas",
            "Targets TRL: MGIT 44 dias",
        ]
        for row, item in enumerate(original, 47):
            ws.merge_cells(start_row=row, start_column=1,
                           end_row=row, end_column=6)
            ws.cell(row, 1, item)

        remove_empty_summary_alerts(wb)
        for row, expected in enumerate(original, 41):
            self.assertEqual(ws.cell(row, 1).value, expected)
        self.assertNotIn("ALERTAS", str(ws["A41"].value))
        self.assertIsNone(ws["A47"].value)

        # Reaplicar ao mesmo ficheiro não deve destruir as notas.
        remove_empty_summary_alerts(wb)
        self.assertEqual(ws["A42"].value, "LEGENDA")

    def test_rejections_export_only_normalized_reasons(self):
        wb = Workbook()
        wb.active.title = "Resumo"
        data = {}
        for lab in LABS:
            data[lab] = {
                "fluxo": [{"Codigo": "Rejeitadas", "Total": 3}],
                "rejeicoes": [
                    {"MotivoAgrupado": "Volume insuficiente",
                     "CodigoOriginal": "HIDDEN_INTERNAL_001",
                     "DescricaoOriginal": "HIDDEN_RAW_TEXT_001",
                     "Quantidade": 2},
                    {"MotivoAgrupado": "Volume insuficiente",
                     "CodigoOriginal": "HIDDEN_INTERNAL_002",
                     "DescricaoOriginal": "HIDDEN_RAW_TEXT_002",
                     "Quantidade": 1},
                ],
            }

        write_rejections(wb, data, date(2026, 9, 13), date(2026, 9, 19))
        ws = wb["Rejeições"]
        vals = [
            str(cell.value)
            for row in ws.iter_rows()
            for cell in row if cell.value is not None
        ]
        self.assertFalse(any("HIDDEN_INTERNAL" in value for value in vals))
        self.assertFalse(any("HIDDEN_RAW_TEXT" in value for value in vals))
        self.assertNotIn("Código original", vals)
        self.assertNotIn("Descrição original", vals)
        self.assertIn("Laboratório", vals)
        self.assertIn("Motivo agrupado", vals)
        self.assertIn("Quantidade", vals)
        self.assertIn("% no laboratório", vals)

        # Um único motivo por laboratório após consolidação das variantes.
        self.assertEqual(sum(
            1 for row in ws.iter_rows()
            if row[0].value == "Carmelo" and row[1].value == "Volume insuficiente"
        ), 1)


if __name__ == "__main__":
    unittest.main()
